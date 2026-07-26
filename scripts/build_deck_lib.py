# -*- coding: utf-8 -*-
"""一品多区方案 PPTX —— 版式与视觉系统辅助库"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 画布 ----------
SW, SH = Inches(13.333), Inches(7.5)
M = Inches(0.62)                      # 左右页边距
CW = SW - 2 * M                       # 内容宽度
BODY_TOP = Inches(1.62)               # 正文起始 y
BODY_BOT = Inches(6.95)               # 正文底线
BODY_H = BODY_BOT - BODY_TOP

# ---------- 色板 ----------
INK      = RGBColor(0x16, 0x20, 0x2E)   # 主文字
INK2     = RGBColor(0x47, 0x57, 0x69)   # 次文字
MUTED    = RGBColor(0x76, 0x85, 0x96)   # 弱文字
LINE     = RGBColor(0xD6, 0xDE, 0xE7)   # 描边
BG_SOFT  = RGBColor(0xF4, 0xF7, 0xFA)   # 浅底
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

BLUE     = RGBColor(0x1D, 0x4E, 0xD8)   # 平台 / WHC
BLUE_BG  = RGBColor(0xEC, 0xF2, 0xFE)
GREEN    = RGBColor(0x04, 0x7A, 0x55)   # 执行 / WMS
GREEN_BG = RGBColor(0xE8, 0xF6, 0xF0)
AMBER    = RGBColor(0xB4, 0x62, 0x05)   # 注意
AMBER_BG = RGBColor(0xFD, 0xF3, 0xE2)
RED      = RGBColor(0xB4, 0x21, 0x21)   # 风险
RED_BG   = RGBColor(0xFD, 0xEC, 0xEC)
PURPLE   = RGBColor(0x5B, 0x30, 0xB8)   # 洞察
PURPLE_BG= RGBColor(0xF2, 0xEE, 0xFD)
SLATE    = RGBColor(0x33, 0x41, 0x55)
SLATE_BG = RGBColor(0xEE, 0xF1, 0xF5)

FONT = "Microsoft YaHei"
FONT_EN = "Segoe UI"


# ---------- 基础 ----------
def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _set_ea(run, name):
    """python-pptx 只设 latin 字体, CJK 需单独写 a:ea"""
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def style(run, size=14, bold=False, color=INK, font=FONT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    _set_ea(run, font)
    return run


def textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def para(tf, first=False, space_before=0, space_after=6, line=1.22, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    p.alignment = align
    return p


def write(tf, chunks, first=False, **kw):
    """chunks: [(text, {style kwargs}), ...] 或 单个 str"""
    p = para(tf, first=first, **{k: v for k, v in kw.items()
                                 if k in ("space_before", "space_after", "line", "align")})
    if isinstance(chunks, str):
        chunks = [(chunks, {})]
    for txt, st in chunks:
        style(p.add_run(), **st).text = txt
    return p


def rect(slide, l, t, w, h, fill=None, line_color=None, line_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=False):
    sp = slide.shapes.add_shape(shape, l, t, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    if not shadow:
        sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.16)
    tf.margin_top = tf.margin_bottom = Inches(0.11)
    return sp


def bar(slide, l, t, w, h, color):
    return rect(slide, l, t, w, h, fill=color, shape=MSO_SHAPE.RECTANGLE)


# ---------- 版式 ----------
def slide_title(slide, title, kicker=None, sub=None):
    """标准页眉:小标签 + 主标题 + 说明"""
    y = Inches(0.46)
    if kicker:
        _, tf = textbox(slide, M, y, CW, Inches(0.24))
        write(tf, [(kicker.upper(), dict(size=10.5, bold=True, color=BLUE, font=FONT_EN))],
              first=True, space_after=0)
        y += Inches(0.30)
    _, tf = textbox(slide, M, y, CW, Inches(0.46))
    write(tf, [(title, dict(size=25, bold=True, color=INK))], first=True, space_after=0)
    y += Inches(0.52)
    if sub:
        _, tf2 = textbox(slide, M, y, CW, Inches(0.3))
        write(tf2, [(sub, dict(size=12.5, color=MUTED))], first=True, space_after=0)
        y += Inches(0.32)
    bar(slide, M, y + Inches(0.06), Inches(0.62), Pt(3), BLUE)
    return y + Inches(0.30)


def footer(slide, idx, note=None):
    _, tf = textbox(slide, M, SH - Inches(0.44), CW, Inches(0.26))
    p = para(tf, first=True, space_after=0)
    if note:
        style(p.add_run(), size=9.5, color=MUTED).text = note
    _, tf2 = textbox(slide, SW - M - Inches(1.0), SH - Inches(0.44), Inches(1.0), Inches(0.26))
    p2 = para(tf2, first=True, space_after=0, align=PP_ALIGN.RIGHT)
    style(p2.add_run(), size=9.5, color=MUTED, font=FONT_EN).text = f"{idx:02d}"


def card(slide, l, t, w, h, title, lines, accent=BLUE, bg=BLUE_BG,
         title_size=13.5, body_size=11, tag=None):
    """带色条的信息卡"""
    sp = rect(slide, l, t, w, h, fill=bg, line_color=None)
    bar(slide, l, t, Pt(3.2), h, accent)
    tf = sp.text_frame
    if tag:
        write(tf, [(tag, dict(size=9.5, bold=True, color=accent, font=FONT_EN))],
              first=True, space_after=3)
        write(tf, [(title, dict(size=title_size, bold=True, color=INK))], space_after=5)
    else:
        write(tf, [(title, dict(size=title_size, bold=True, color=INK))],
              first=True, space_after=5)
    for ln in lines:
        if isinstance(ln, str):
            write(tf, [("· ", dict(size=body_size, color=accent, bold=True)),
                       (ln, dict(size=body_size, color=INK2))], space_after=3, line=1.2)
        elif isinstance(ln, tuple) and len(ln) == 2 and isinstance(ln[1], dict):
            # 单个 (文本, 样式) 组 —— 包成 chunk 列表
            write(tf, [ln], space_after=3, line=1.2)
        else:
            write(tf, ln, space_after=3, line=1.2)
    return sp


def kv_table(slide, l, t, w, headers, rows, col_ratio, row_h=Inches(0.42),
             head_h=Inches(0.40), font_size=10.5, head_size=10.5, zebra=True,
             cell_colors=None):
    """cell_colors: {(r,c): (bg, fg, bold)}"""
    n_rows, n_cols = len(rows) + 1, len(headers)
    total_h = head_h + row_h * len(rows)
    gt = slide.shapes.add_table(n_rows, n_cols, l, t, w, total_h).table
    gt.first_row = True
    gt.horz_banding = False

    tot = sum(col_ratio)
    for i, r in enumerate(col_ratio):
        gt.columns[i].width = Emu(int(w * r / tot))
    gt.rows[0].height = head_h
    for i in range(1, n_rows):
        gt.rows[i].height = row_h

    for c, htxt in enumerate(headers):
        cell = gt.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = SLATE
        cell.margin_left = cell.margin_right = Inches(0.10)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        write(tf, [(htxt, dict(size=head_size, bold=True, color=WHITE))], first=True, space_after=0)

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = gt.cell(r, c)
            bgc = BG_SOFT if (zebra and r % 2 == 0) else WHITE
            fgc, bold = INK2, False
            if cell_colors and (r - 1, c) in cell_colors:
                cc = cell_colors[(r - 1, c)]
                bgc = cc[0] if cc[0] is not None else bgc
                fgc = cc[1] if len(cc) > 1 and cc[1] is not None else fgc
                bold = cc[2] if len(cc) > 2 else False
            cell.fill.solid()
            cell.fill.fore_color.rgb = bgc
            cell.margin_left = cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            write(tf, [(str(val), dict(size=font_size, color=fgc, bold=bold))],
                  first=True, space_after=0, line=1.15)
    return gt


def chevron(slide, l, t, w, h, num, title, desc, accent, bg):
    """流程步骤块"""
    sp = rect(slide, l, t, w, h, fill=bg, line_color=None)
    tf = sp.text_frame
    write(tf, [(f"STEP {num}", dict(size=9, bold=True, color=accent, font=FONT_EN))],
          first=True, space_after=3)
    write(tf, [(title, dict(size=12, bold=True, color=INK))], space_after=4)
    write(tf, [(desc, dict(size=10, color=INK2))], space_after=0, line=1.18)
    return sp


def arrow_right(slide, l, t, w, h, color=LINE):
    sp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, l, t, w, h)
    sp.rotation = 90
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def section_divider(prs, idx, num, title, sub, points):
    s = blank(prs)
    bar(s, Emu(0), Emu(0), SW, SH, INK)
    _, tf = textbox(s, M + Inches(0.3), Inches(2.35), Inches(6.4), Inches(0.4))
    write(tf, [(f"PART {num}", dict(size=12, bold=True, color=RGBColor(0x7E, 0xA6, 0xF5), font=FONT_EN))],
          first=True, space_after=8)
    _, tf2 = textbox(s, M + Inches(0.3), Inches(2.85), Inches(6.6), Inches(0.7))
    write(tf2, [(title, dict(size=34, bold=True, color=WHITE))], first=True, space_after=0)
    _, tf3 = textbox(s, M + Inches(0.3), Inches(3.72), Inches(6.2), Inches(0.5))
    write(tf3, [(sub, dict(size=13, color=RGBColor(0xA8, 0xB6, 0xC8)))], first=True, space_after=0)
    bar(s, M + Inches(0.3), Inches(4.32), Inches(0.9), Pt(3), RGBColor(0x7E, 0xA6, 0xF5))

    x = Inches(7.9)
    y = Inches(2.35)
    for pt in points:
        _, ptf = textbox(s, x, y, Inches(4.6), Inches(0.5))
        write(ptf, [("— ", dict(size=12, color=RGBColor(0x7E, 0xA6, 0xF5), bold=True)),
                    (pt, dict(size=12.5, color=RGBColor(0xD3, 0xDC, 0xE7)))],
              first=True, space_after=0, line=1.3)
        y += Inches(0.62)
    return s
